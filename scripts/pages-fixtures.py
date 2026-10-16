"""Generate small real document fixtures; never used by the published site."""
import sys
from pathlib import Path
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

folder = Path(sys.argv[1])
text = 'La sensibilita rimane conservata nella miastenia gravis. La debolezza fluttuante richiede una valutazione neurologica accurata.'
doc = Document();doc.add_heading('Capitolo di neurologia', 1);doc.add_paragraph(text);doc.save(folder/'appunti.docx')
ppt = Presentation();slide = ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='Neurologia';slide.placeholders[1].text=text;ppt.save(folder/'appunti.pptx')
writer = PdfWriter();page = writer.add_blank_page(width=600, height=800)
font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
stream = DecodedStreamObject();stream.set_data(('BT /F1 10 Tf 20 700 Td ('+text+') Tj ET').encode('ascii'))
page[NameObject('/Contents')] = writer._add_object(stream)
with (folder/'appunti.pdf').open('wb') as f: writer.write(f)
blank=PdfWriter();blank.add_blank_page(width=600,height=800)
with (folder/'scansione.pdf').open('wb') as f: blank.write(f)
