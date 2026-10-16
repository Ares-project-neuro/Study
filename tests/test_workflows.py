import io
from datetime import datetime, timedelta, timezone
import pytest
from fastapi.testclient import TestClient
from app.main import app, LOGIN_ATTEMPTS
from app import store, ai


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('STUDY_DATA_DIR', str(tmp_path))
    monkeypatch.delenv('STUDY_AI_KEY', raising=False)
    monkeypatch.delenv('STUDY_PASSWORD', raising=False)
    LOGIN_ATTEMPTS.clear()
    with TestClient(app) as c:
        c.get('/api/status')
        assert c.post('/api/login', json={'password': 'only-tests-password-123'}).status_code == 200
        yield c


TEXT = ('La miastenia gravis presenta debolezza fluttuante e affaticabilità muscolare. '
        'La sensibilità generalmente rimane conservata nel disturbo della giunzione neuromuscolare. '
        'La diagnosi richiede integrazione della storia clinica e degli accertamenti neurofisiologici.')


def upload(client, content=TEXT, name='Appunti.txt', **fields):
    r = client.post('/api/documents', files={'file': (name, content.encode())}, data={'topic': 'Giunzione', **fields})
    assert r.status_code == 200, r.text
    return r.json()


def seed(client):
    assert client.post('/api/demo').status_code == 200


def test_auth_csrf_and_export_privacy(client):
    assert client.post('/api/demo', headers={'Origin': 'https://evil.invalid'}).status_code == 403
    exported = client.get('/api/export').json()
    assert 'auth' not in exported and 'signing' not in str(exported)
    client.post('/api/logout')
    assert client.get('/api/documents').status_code == 401
    assert client.post('/api/login', json={'password': 'wrong-password-123'}).status_code == 401


def test_upload_generate_quiz_and_persist(client):
    doc = upload(client)
    generated = client.post('/api/generate', json={'document_ids': [doc['id']], 'count': 2}).json()
    assert len(generated['items']) == 2
    for q in generated['items']:
        assert q['sources'][0]['quote'] in doc['sections'][0]['text']
        assert q['review_required']
    started = client.post('/api/quizzes', json={'include_review': True, 'count': 2}).json()
    assert len(started['questions']) == 2
    q = started['questions'][0]
    assert 'correct' not in q and 'sources' not in q
    correct = next(i for i in generated['items'] if i['id'] == q['id'])['correct']
    answered = client.post(f'/api/quizzes/{started["id"]}/answer', json={'question_id': q['id'], 'selected': correct}).json()
    assert answered['answers'][q['id']]['correct']
    assert answered['questions'][0]['explanation']
    resumed = client.get(f'/api/quizzes/{started["id"]}').json()
    assert resumed['answers'] == answered['answers']
    assert client.post(f'/api/quizzes/{started["id"]}/answer', json={'question_id': q['id'], 'selected': correct}).status_code == 409
    finished = client.post(f'/api/quizzes/{started["id"]}/finish').json()
    assert all('correct' in q for q in finished['questions'])
    # Connection reopened: persistent rows, not process-local state.
    assert store.get('quiz', started['id'])['finished_at']


def test_exam_no_answer_leak(client):
    seed(client)
    quiz = client.post('/api/quizzes', json={'include_demo': True, 'mode': 'exam'}).json()
    q = quiz['questions'][0]
    result = client.post(f'/api/quizzes/{quiz["id"]}/answer', json={'question_id': q['id'], 'selected': [0]}).json()
    assert 'correct' not in result['answers'][q['id']]
    assert all('correct' not in question and 'explanation' not in question for question in result['questions'])
    exported = client.get('/api/export').json()['quizzes'][0]
    assert 'correct' not in exported['answers'][q['id']]
    assert client.post(f'/api/quizzes/{quiz["id"]}/finish').json()['questions'][0]['correct']


def test_quiz_single_and_multiple_validation(client):
    seed(client)
    quiz = client.post('/api/quizzes', json={'include_demo': True}).json()
    single = next(q for q in quiz['questions'] if q['type'] == 'single')
    multi = next(q for q in quiz['questions'] if q['type'] == 'multiple')
    path = f'/api/quizzes/{quiz["id"]}/answer'
    assert client.post(path, json={'question_id': single['id'], 'selected': [0, 1]}).status_code == 422
    assert client.post(path, json={'question_id': multi['id'], 'selected': [0, 1]}).json()['answers'][multi['id']]['correct']


def test_source_versions_and_explicit_regeneration(client):
    d = upload(client)
    q = client.post('/api/generate', json={'document_ids': [d['id']], 'count': 1}).json()['items'][0]
    updated = upload(client, document_id=d['id'], mark_linked='true')
    assert updated['revision'] == 2
    old = store.get('question', q['id'])
    assert old['sources'][0]['revision'] == 1
    assert 'aggiornata' in old['review_reason']
    replacement = client.post('/api/generate', json={'document_ids': [d['id']], 'count': 1, 'replace_id': q['id']}).json()
    assert len(replacement['items']) == 1
    assert replacement['items'][0]['sources'][0]['revision'] == 2
    client.delete('/api/documents/' + d['id'])
    assert 'eliminata' in store.get('question', q['id'])['review_reason']


def test_section_selection_and_duplicate_rejection(client):
    d = upload(client, TEXT + '\n\n' + TEXT.replace('miastenia', 'patologia'))
    assert len(d['sections']) == 2
    body = {'document_ids': [d['id']], 'sections': {d['id']: ['2']}, 'kind': 'card', 'count': 1}
    first = client.post('/api/generate', json=body).json()
    assert first['items'][0]['sources'][0]['section_id'] == '2'
    assert client.post('/api/generate', json=body).json()['items'] == []


@pytest.mark.parametrize('grade,min_days,max_days', [('again', .006, .008), ('hard', .99, 1.01), ('good', .99, 1.01), ('easy', 3.99, 4.01)])
def test_spaced_repetition_persists(client, grade, min_days, max_days):
    card = client.post('/api/cards', json={'front': 'Dove si trova il disturbo?', 'back': 'Giunzione neuromuscolare'}).json()
    review = client.post('/api/reviews', json={}).json()
    assert 'back' not in review['cards'][0]
    path = f'/api/reviews/{review["id"]}'
    assert client.post(path+'/rate', json={'card_id': card['id'], 'grade': grade}).status_code == 409
    assert client.post(path+'/reveal/'+card['id']).json()['cards'][0]['back'] == card['back']
    result = client.post(path+'/rate', json={'card_id': card['id'], 'grade': grade}).json()
    due = datetime.fromisoformat(result['ratings'][card['id']]['due'])
    delta = (due-datetime.now(timezone.utc)).total_seconds()/86400
    assert min_days < delta < max_days
    saved = store.get('card', card['id'])
    assert saved['reviews'] == 1 and saved['due'] == due.isoformat()
    assert client.post(path+'/rate', json={'card_id': card['id'], 'grade': grade}).status_code == 409


def test_cloze_hidden_and_daily_load(client):
    client.put('/api/settings', json={'daily_limit': 1, 'daily_goal': 1})
    client.post('/api/cards', json={'front': 'La sensibilità è {{conservata}}', 'back': 'conservata', 'type': 'cloze'})
    client.post('/api/cards', json={'front': 'Altra domanda?', 'back': 'Altra risposta'})
    review = client.post('/api/reviews', json={}).json()
    assert len(review['cards']) == 1
    c = review['cards'][0]
    assert 'back' not in c and '{{' not in c['front']
    client.post(f'/api/reviews/{review["id"]}/reveal/{c["id"]}')
    client.post(f'/api/reviews/{review["id"]}/rate', json={'card_id': c['id'], 'grade': 'good'})
    assert client.post('/api/reviews', json={}).status_code == 422


def test_case_consistency_privacy_and_ambiguous_requests(client):
    case = client.post('/api/cases', json={'topic': 'Malattie neuromuscolari'}).json()
    path = '/api/cases/' + case['id']
    assert 'dossier' not in case and 'diagnosis' not in case
    assert 'reports' not in str(case)
    first = client.post(path+'/actions', json={'action': 'history'}).json()['events'][-1]['response']
    repeated = client.post(path+'/actions', json={'text': 'Da quanto tempo sono presenti i sintomi?'}).json()['events'][-1]['response']
    assert first == repeated
    hypothesis = client.post(path+'/actions', json={'text': 'Ipotizzo miastenia'}).json()['events'][-1]
    assert hypothesis['kind'] == 'hypothesis' and 'action' not in hypothesis
    ambiguous = client.post(path+'/actions', json={'text': 'Fai imaging'}).json()['events'][-1]
    assert ambiguous['kind'] == 'clarify' and 'action' not in ambiguous
    client.put(path+'/notes', json={'notes': 'Giunzione neuromuscolare da valutare'})
    assert client.get(path).json()['notes'].startswith('Giunzione')
    exported = client.get('/api/export').json()['cases'][0]
    assert 'dossier' not in exported and 'assessment' not in exported
    assert client.post(path+'/hint').status_code == 200


def test_grading_uses_unique_actions_and_guess_is_not_enough(client):
    opts = {'topic': 'Malattie neuromuscolari', 'mode': 'exam'}
    case = client.post('/api/cases', json=opts).json()
    path = '/api/cases/' + case['id']
    assert client.post(path+'/hint').status_code == 409
    for _ in range(5):
        client.post(path+'/actions', json={'action': 'history'})
    result = client.post(path+'/finish', json={'diagnosis': 'miastenia', 'rationale': ''}).json()
    assert result['assessment']['total'] < 40
    assert result['assessment']['dimensions'][0]['score'] < 10
    assert client.post(path+'/actions', json={'action': 'neuro'}).status_code == 409
    assert client.post(path+'/finish', json={'diagnosis': 'altro'}).json()['submission']['diagnosis'] == 'miastenia'


def test_complete_case_scored_path_and_reports(client):
    c = client.post('/api/cases', json={'topic': 'Malattie neuromuscolari', 'level': 'specializzando'}).json()
    p = '/api/cases/' + c['id']
    for action in ['history', 'respiration', 'neuro', 'antibodies', 'emg']:
        assert client.post(p+'/actions', json={'action': action}).status_code == 200
    result = client.post(p+'/finish', json={'diagnosis': 'miastenia gravis', 'rationale': 'Affaticabilità fluttuante con ptosi e decremento alla stimolazione.', 'localization': 'Giunzione neuromuscolare postsinaptica', 'alternatives': 'Lambert-Eaton, miopatia'}).json()
    assert result['assessment']['stars'] == 5
    assert result['assessment']['omissions'] == []
    assert client.post('/api/reports', json={'kind': 'case', 'target_id': c['id'], 'text': 'Vorrei rivalutare un percorso alternativo'}).status_code == 200
    assert client.get('/api/export').json()['reports']


def test_deadlines_are_enforced(client):
    seed(client)
    q = client.post('/api/quizzes', json={'include_demo': True, 'minutes': 1}).json()
    saved = store.get('quiz', q['id'])
    saved['deadline'] = (datetime.now(timezone.utc)-timedelta(seconds=1)).isoformat()
    store.save('quiz', saved)
    assert client.post(f'/api/quizzes/{q["id"]}/answer', json={'question_id': q['questions'][0]['id'], 'selected': [0]}).status_code == 409
    assert client.post(f'/api/quizzes/{q["id"]}/finish').status_code == 200


def test_office_and_unreadable_extraction(client):
    from docx import Document
    from pptx import Presentation
    doc = Document()
    doc.add_heading('Capitolo', 1)
    doc.add_paragraph(TEXT)
    buffer = io.BytesIO();doc.save(buffer)
    r = client.post('/api/documents', files={'file': ('test.docx', buffer.getvalue())})
    assert r.status_code == 200 and 'miastenia' in r.json()['sections'][0]['text']
    ppt = Presentation();slide = ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text = 'Neurologia';slide.placeholders[1].text = TEXT
    buffer = io.BytesIO();ppt.save(buffer)
    r = client.post('/api/documents', files={'file': ('test.pptx', buffer.getvalue())})
    assert r.status_code == 200 and r.json()['sections'][0]['label'] == 'Diapositiva 1'
    from pypdf import PdfWriter
    pdf = PdfWriter();pdf.add_blank_page(width=100, height=100);buffer = io.BytesIO();pdf.write(buffer)
    r = client.post('/api/documents', files={'file': ('scan.pdf', buffer.getvalue())})
    assert r.status_code == 422 and 'OCR' in r.json()['detail']
    assert client.post('/api/documents', files={'file': ('bad.docx', b'not zip')}).status_code == 422


def test_ai_unavailable_and_bad_citations_rejected(client, monkeypatch):
    doc = upload(client)
    body = {'document_ids': [doc['id']], 'engine': 'ai'}
    assert client.post('/api/generate', json=body).status_code == 409
    monkeypatch.setattr(ai, 'complete', lambda *args: {'items': [{'source_key': doc['id']+':1', 'quote': 'Una citazione completamente inventata', 'prompt': 'Non valido'}], 'warnings': []})
    result = client.post('/api/generate', json=body).json()
    assert result['items'] == [] and any('citazione' in s for s in result['warnings'])


def test_demo_separation_and_idempotency(client):
    seed(client);seed(client)
    assert len(client.get('/api/documents').json()) == 1
    assert client.post('/api/quizzes', json={}).status_code == 422
    assert client.post('/api/reviews', json={}).status_code == 422
    dash = client.get('/api/dashboard').json()
    assert dash['metrics'] == {} and dash['due'] == 0


def test_grader_does_not_reward_uncollected_evidence_or_negated_diagnosis(client):
    c = client.post('/api/cases', json={'topic': 'Malattie neuromuscolari'}).json()
    result = client.post('/api/cases/'+c['id']+'/finish', json={'diagnosis': 'Escludo miastenia', 'rationale': 'Ptosi fluttuante e decremento agli esami non ancora richiesti'}).json()
    assert not result['assessment']['correct']
    assert not result['assessment']['evidence_terms']


def test_semantic_judge_rejects_invented_events(client, monkeypatch):
    from app import cases
    c = client.post('/api/cases', json={'topic': 'Malattie neuromuscolari'}).json()
    private = store.get('case', c['id'])
    submission = {'diagnosis': 'miastenia', 'rationale': '', 'localization': '', 'alternatives': ''}
    fallback = cases.evaluate(private, submission)
    response = {'dimensions': [{'key': d['key'], 'score': d['weight'], 'reason': 'Ottimo', 'event_indices': [999]} for d in cases.RUBRIC],
                'strengths': [], 'omissions': [], 'unnecessary': [], 'correct': True, 'explanation': 'Dati inventati'}
    monkeypatch.setattr(ai, 'complete', lambda *args: response)
    result = cases.semantic_evaluation(private, submission, fallback)
    assert result['engine'] == 'euristico locale' and result['total'] == fallback['total']
    assert 'non ha restituito' in result['limits']


def test_semantic_judge_preserves_weights_and_stars(client, monkeypatch):
    from app import cases
    c = client.post('/api/cases', json={'topic': 'Malattie neuromuscolari'}).json()
    client.post('/api/cases/'+c['id']+'/actions', json={'action': 'history'})
    private = store.get('case', c['id'])
    submission = {'diagnosis': 'miastenia', 'rationale': '', 'localization': '', 'alternatives': ''}
    fallback = cases.evaluate(private, submission)
    response = {'dimensions': [{'key': d['key'], 'score': 5, 'reason': 'Evidenza parziale nel primo evento', 'event_indices': [0]} for d in cases.RUBRIC],
                'strengths': [{'event_index': 0, 'reason': 'Esordio raccolto'}], 'omissions': [{'action_id': 'respiration', 'reason': 'Sicurezza respiratoria non valutata'}],
                'unnecessary': [], 'correct': True, 'explanation': 'Punteggio parziale: percorso incompleto'}
    monkeypatch.setattr(ai, 'complete', lambda *args: response)
    result = cases.semantic_evaluation(private, submission, fallback)
    assert result['engine'] == 'IA semantica' and result['total'] == 20 and result['stars'] == 1
    assert [d['weight'] for d in result['dimensions']] == [20, 25, 25, 30]


def test_ai_valid_generation_has_verified_source_and_review_status(client, monkeypatch):
    d = upload(client)
    item = {'source_key': d['id']+':1', 'quote': TEXT.split('. ')[0]+'.', 'topic': 'Wrong model topic',
            'prompt': 'Quale caratteristica è riportata nel materiale caricato?', 'type': 'single',
            'options': ['Affaticabilità', 'Anestesia completa', 'Debolezza invariabile', 'Solo sintomi sensitivi'], 'correct': [0],
            'explanation': 'Il materiale descrive affaticabilità.', 'option_explanations': ['Presente', 'Non supportata', 'Non supportata', 'Non supportata']}
    monkeypatch.setattr(ai, 'complete', lambda *args: {'items': [item], 'warnings': ['Verifica clinica necessaria']})
    r = client.post('/api/generate', json={'document_ids': [d['id']], 'engine': 'ai', 'count': 1}).json()
    assert len(r['items']) == 1 and r['items'][0]['review_required']
    assert r['items'][0]['topic'] == d['topic'] and r['items'][0]['origin'] == 'ai'
