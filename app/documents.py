import io
import re
import zipfile
from fastapi import HTTPException

MAX_BYTES = 20 * 1024 * 1024


def extract(name, content):
    """Return individually addressable pages/sections, keeping partial failures visible."""
    ext = name.lower().rsplit('.', 1)[-1]
    sections, warnings = [], []

    def add(label, text):
        text = re.sub(r'[ \t]+', ' ', text or '').strip()
        if text:
            sections.append({'id': str(len(sections) + 1), 'label': label, 'text': text})
        else:
            warnings.append(f'{label}: nessun testo estratto; potrebbe essere necessario OCR.')

    if ext in ('docx', 'pptx'):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                    raise HTTPException(413, 'Documento decompresso troppo grande (massimo 80 MB)')
        except zipfile.BadZipFile:
            raise HTTPException(422, 'Il file non è un documento Office valido')
    try:
        if ext == 'pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise HTTPException(422, 'PDF protetto: carica una copia leggibile senza password')
            if len(reader.pages) > 600:
                raise HTTPException(413, 'Massimo 600 pagine per documento')
            for i, page in enumerate(reader.pages):
                try:
                    add(f'Pagina {i+1}', page.extract_text())
                except Exception:
                    warnings.append(f'Pagina {i+1}: estrazione fallita, parte non elaborata.')
        elif ext == 'docx':
            from docx import Document
            doc = Document(io.BytesIO(content))
            label, paragraphs = 'Sezione 1', []
            for p in doc.paragraphs:
                if p.style.name.startswith('Heading') and paragraphs:
                    add(label, '\n'.join(paragraphs))
                    label, paragraphs = p.text, []
                paragraphs.append(p.text)
            if paragraphs:
                add(label, '\n'.join(paragraphs))
            for i, table in enumerate(doc.tables):
                add(f'Tabella {i+1}', '\n'.join(' | '.join(c.text for c in r.cells) for r in table.rows))
            warnings.append('DOCX: numeri di pagina non disponibili; immagini, note, intestazioni e piè di pagina non elaborati.')
        elif ext == 'pptx':
            from pptx import Presentation
            for i, slide in enumerate(Presentation(io.BytesIO(content)).slides):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
                    if shape.has_table:
                        texts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                add(f'Diapositiva {i+1}', '\n'.join(texts))
            warnings.append('PPTX: immagini, grafici, note relatore e gruppi di forme non elaborati.')
        elif ext in ('txt', 'md'):
            text = content.decode('utf-8-sig')
            for i, part in enumerate(re.split(r'\n\s*\n', text)):
                if part.strip():
                    add(f'Paragrafo {i+1}', part)
        else:
            raise HTTPException(415, 'Formati ammessi: PDF, DOCX, PPTX, TXT UTF-8 e MD')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Documento illeggibile o formato non valido. Per scansioni usa OCR prima del caricamento.')
    if not sections:
        raise HTTPException(422, 'Nessun testo leggibile. Il documento può richiedere OCR; nessun contenuto è stato generato.')
    if sum(len(s['text']) for s in sections) > 1_500_000:
        raise HTTPException(413, 'Troppo testo: suddividi il documento in parti più piccole')
    if ext == 'pdf':
        warnings.append('PDF: immagini e formule non vengono interpretate; verifica ordine del testo e tabelle.')
    return sections, warnings
